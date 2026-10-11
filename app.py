"""
z/OS 3.1 -> 3.2 Migration Fix Dashboard
---------------------------------------
Streamlit app that:
  * keeps an inventory of PTFs, USERMODs, APARs, FIXCATs and other fixes
  * ingests your own SMP/E output (REPORT MISSINGFIX, LIST, HOLDDATA, CSI exports)
    and any IBM web page / pasted text, extracting fix IDs automatically
  * offers a RAG assistant (TF-IDF retrieval via scikit-learn, optional Claude answers)
  * clusters fixes with ML (TF-IDF -> SVD -> KMeans) to group related work
  * exports everything to a formatted Excel workbook
  * lists the IBM sites where the authoritative details live

Run:  streamlit run app.py
"""

import io
import os
import re
from datetime import datetime

import numpy as np
import pandas as pd
import plotly.express as px
import requests
import streamlit as st
from bs4 import BeautifulSoup
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sklearn.cluster import KMeans
from sklearn.decomposition import TruncatedSVD
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

st.set_page_config(page_title="z/OS 3.2 Migration Fix Dashboard", page_icon="🖥️", layout="wide")

CLAUDE_MODEL = "claude-sonnet-5-5"

# --------------------------------------------------------------------------------------
# Reference data
# --------------------------------------------------------------------------------------
# Authoritative places to get z/OS 3.2 product / service details.
# Deep links change often, so entry points are listed; use the "Search hint" column.
SOURCES = pd.DataFrame(
    [
        ["IBM z/OS product page", "https://www.ibm.com/products/zos", "Product overview, announcements, what's new in 3.2", "Planning"],
        ["IBM Documentation - z/OS", "https://www.ibm.com/docs/en/zos", "z/OS 3.2 library: Migration, Release Planning, Upgrade Workflow", "Migration"],
        ["z/OS Upgrade Workflow (in IBM Docs / z/OSMF)", "https://www.ibm.com/docs/en/zos", "Step-by-step 3.1 -> 3.2 upgrade actions; search 'z/OS Upgrade Workflow'", "Migration"],
        ["z/OS Migration guide (MVS, JES2/3, Security, etc.)", "https://www.ibm.com/docs/en/zos", "Per-element migration actions; search 'Migration from z/OS V3R1 to V3R2'", "Migration"],
        ["ShopZ (ServerPac, CBPDO, SMP/E orders)", "https://www.ibm.com/software/shopzseries/ShopzSeries_public.wss", "Order z/OS 3.2, PTFs, FIXCAT-based service, HOLDDATA", "Ordering"],
        ["IBM Fix Central", "https://www.ibm.com/support/fixcentral/", "Download fixes and service for IBM products", "Fixes"],
        ["IBM Support - z/OS", "https://www.ibm.com/support/home/product/D367161S68764U34/other_software/z_os", "Support home, technical documents, APAR search", "Support"],
        ["Preventive Service Planning (PSP) buckets", "https://www.ibm.com/support/pages/", "Search 'PSP bucket ZOSV3R2' / 'ZOSV3R1' for upgrade-specific HIPER/PE/FIXCAT info", "PSP"],
        ["Recommended Service Upgrade (RSU)", "https://www.ibm.com/support/pages/", "Search 'RSU' for current z/OS recommended service level", "Service"],
        ["Fix Categories (FIXCAT)", "https://www.ibm.com/support/pages/", "Search 'IBM Fix Category Values and Descriptions'", "Fixes"],
        ["IBM Z Software Support - APAR / PTF search", "https://www.ibm.com/support/pages/", "Search by APAR (OAxxxxx/PHxxxxx), PTF (UJxxxxx/UAxxxxx) or component", "APAR"],
        ["IBM Z HOLDDATA", "https://service.boulder.ibm.com/390holddata.html", "Latest HOLDDATA (ERROR/FIXCAT/SYSTEM/USER HOLDs) for SMP/E", "HOLDDATA"],
        ["IBM Redbooks", "https://www.redbooks.ibm.com/", "z/OS 3.2 migration / implementation guides", "Learning"],
        ["IBM Z and LinuxONE Community", "https://community.ibm.com/community/user/ibmz-and-linuxone/home", "Announcements, discussions, migration experience", "Community"],
        ["SHARE Association", "https://www.share.org/", "z/OS migration sessions and handouts", "Community"],
        ["z/OS Preventive Service Planning (IBM Z Software Support)", "https://www.ibm.com/support/pages/", "Search 'z/OS V3R2 PSP' for subset-level service", "PSP"],
    ],
    columns=["Source", "URL", "What you get / search hint", "Category"],
)

# REPORT MISSINGFIX is driven by the FIXCAT HOLDs in the HOLDDATA you RECEIVE, so the
# same IBM sources feed both. Tag them so the dashboard shows what each site is used for.
SOURCES["Feeds"] = "Reference"
_feeds_both = SOURCES["URL"].str.contains("390holddata|shopzseries|fixcentral", case=False)
SOURCES.loc[_feeds_both, "Feeds"] = "HOLDDATA + MISSINGFIX"
SOURCES.loc[SOURCES["Category"].isin(["PSP", "Fixes", "APAR", "Service"]) & ~_feeds_both, "Feeds"] = "PSP / FIXCAT / APAR lookup"

# Starter inventory: categories of fixes to check. These are NOT specific PTF numbers -
# real PTF/APAR numbers must come from your HOLDDATA / REPORT MISSINGFIX / PSP output.
SEED = pd.DataFrame(
    [
        ["FIXCAT", "IBM.Coexistence.z/OS.V3R2", "z/OS", "Coexistence fixes required on 3.1 while 3.2 runs in the same sysplex", "Pre-migration (3.1)", "High"],
        ["FIXCAT", "IBM.Migration.z/OS.V3R2", "z/OS", "Fixes IBM recommends before/with migration to 3.2", "Pre-migration (3.1)", "High"],
        ["FIXCAT", "IBM.TargetSystem-RequiredService.z/OS.V3R2", "z/OS", "Service required on the 3.2 target system", "Post-migration (3.2)", "High"],
        ["FIXCAT", "IBM.Device.Server.z17-9175 (and relevant hardware FIXCATs)", "z/OS", "Hardware-support fixes for your CPC model", "Pre-migration (3.1)", "Medium"],
        ["PTF", "(from REPORT MISSINGFIX)", "SMP/E", "Run REPORT MISSINGFIX against each target zone with the FIXCATs above", "Pre-migration (3.1)", "High"],
        ["APAR", "(from PSP bucket ZOSV3R2)", "PSP", "HIPER / PE / special-attention APARs listed in the 3.2 PSP upgrade", "Pre-migration (3.1)", "High"],
        ["USERMOD", "(your site's USERMODs)", "SMP/E", "Review every local USERMOD for supersede / rework / reapply on 3.2", "Post-migration (3.2)", "High"],
        ["OTHER", "SYSTEM/ERROR HOLDs", "HOLDDATA", "Resolve ERROR and SYSTEM HOLDs before APPLY; document bypass reasons", "Pre-migration (3.1)", "Medium"],
    ],
    columns=["Type", "ID", "Component/FMID", "Description", "Phase", "Priority"],
)
SEED["Source"] = "Starter checklist"
SEED["Verified"] = False

# --------------------------------------------------------------------------------------
# Parsing helpers (fix ID extraction from SMP/E output or web pages)
# --------------------------------------------------------------------------------------
RE_APAR = re.compile(r"\b(?:OA|OW|OY|PH|PI|PK|PM|PQ|PJ)\d{5}\b")
RE_PTF = re.compile(r"\bU[A-Z]\d{5}\b")
RE_FIXCAT = re.compile(r"\bIBM\.[A-Za-z0-9_.\-/]+")
RE_HOLD = re.compile(r"\+\+HOLD\(\s*([A-Z0-9]{7})\s*\).*?\b(ERROR|SYSTEM|USER|FIXCAT|EXCLUDE)\b", re.I)
RE_FMID = re.compile(r"\bH[A-Z0-9]{2}\d{4}\b|\bJ[A-Z0-9]{2}\d{4}\b")


def extract_fixes(text: str, source: str, phase: str, usermod_regex: str) -> pd.DataFrame:
    """Scan free text line by line and classify IDs."""
    rows = []
    um_re = re.compile(usermod_regex) if usermod_regex.strip() else None
    for line in text.splitlines():
        line_s = line.strip()
        if not line_s:
            continue
        fmid = RE_FMID.search(line_s)
        fmid = fmid.group(0) if fmid else ""
        found = set()
        for m in RE_FIXCAT.findall(line_s):
            found.add(("FIXCAT", m.rstrip(".,;")))
        # HOLDDATA statements, e.g. ++HOLD(UJ12345) FIXCAT FMID(HBB77D0) REASON(OA12345) ...
        hm = RE_HOLD.search(line_s)
        if hm:
            found.add(("HOLD", f"{hm.group(1).upper()} [{hm.group(2).upper()}]"))
        for m in RE_PTF.findall(line_s):
            found.add(("PTF", m))
        for m in RE_APAR.findall(line_s):
            found.add(("APAR", m))
        if um_re:
            for m in um_re.findall(line_s):
                found.add(("USERMOD", m if isinstance(m, str) else m[0]))
        for typ, fid in found:
            rows.append([typ, fid, fmid, line_s[:300], phase, "Medium", source, True])
    return pd.DataFrame(rows, columns=["Type", "ID", "Component/FMID", "Description", "Phase", "Priority", "Source", "Verified"])


def read_upload(f) -> str:
    name = f.name.lower()
    if name.endswith(".csv"):
        return pd.read_csv(f, dtype=str).fillna("").to_csv(sep=" ", index=False)
    if name.endswith((".xlsx", ".xls")):
        sheets = pd.read_excel(f, sheet_name=None, dtype=str)
        return "\n".join(df.fillna("").to_csv(sep=" ", index=False) for df in sheets.values())
    return f.read().decode("utf-8", errors="ignore")


def fetch_url_text(url: str) -> str:
    r = requests.get(url, timeout=20, headers={"User-Agent": "Mozilla/5.0 (zOS-migration-dashboard)"})
    r.raise_for_status()
    soup = BeautifulSoup(r.text, "html.parser")
    for t in soup(["script", "style", "noscript"]):
        t.decompose()
    return soup.get_text("\n")


# --------------------------------------------------------------------------------------
# State
# --------------------------------------------------------------------------------------
if "inventory" not in st.session_state:
    st.session_state.inventory = SEED.copy()
if "docs" not in st.session_state:
    st.session_state.docs = []  # list of dicts: {"title","text"} for RAG


def add_rows(df: pd.DataFrame):
    if df.empty:
        return 0
    inv = pd.concat([st.session_state.inventory, df], ignore_index=True)
    inv = inv.drop_duplicates(subset=["Type", "ID", "Phase"], keep="last").reset_index(drop=True)
    added = len(inv) - len(st.session_state.inventory)
    st.session_state.inventory = inv
    return max(added, 0)


def add_doc(title: str, text: str, chunk_words: int = 180):
    words = text.split()
    for i in range(0, len(words), chunk_words):
        st.session_state.docs.append({"title": title, "text": " ".join(words[i : i + chunk_words])})


# --------------------------------------------------------------------------------------
# RAG (TF-IDF retrieval; optional Claude generation)
# --------------------------------------------------------------------------------------
def build_corpus():
    corpus = []
    for _, r in st.session_state.inventory.iterrows():
        corpus.append({"title": f"{r['Type']} {r['ID']}", "text": f"{r['Type']} {r['ID']} {r['Component/FMID']} {r['Description']} {r['Phase']} {r['Source']}"})
    for _, r in SOURCES.iterrows():
        corpus.append({"title": f"Source: {r['Source']}", "text": f"{r['Source']} {r['What you get / search hint']} {r['URL']}"})
    corpus.extend(st.session_state.docs)
    return corpus


def retrieve(query: str, k: int = 6):
    corpus = build_corpus()
    texts = [c["text"] for c in corpus]
    vec = TfidfVectorizer(stop_words="english", ngram_range=(1, 2), sublinear_tf=True)
    mat = vec.fit_transform(texts + [query])
    sims = cosine_similarity(mat[-1], mat[:-1]).ravel()
    idx = np.argsort(-sims)[:k]
    return [(corpus[i], float(sims[i])) for i in idx if sims[i] > 0]


def get_api_key() -> str:
    key = st.session_state.get("api_key_input", "").strip()
    if key:
        return key
    try:
        if "ANTHROPIC_API_KEY" in st.secrets:
            return st.secrets["ANTHROPIC_API_KEY"]
    except Exception:
        pass
    return os.environ.get("ANTHROPIC_API_KEY", "")


def ask_claude(question: str, hits) -> str:
    import anthropic

    ctx = "\n\n".join(f"[{i+1}] {h['title']}\n{h['text']}" for i, (h, _) in enumerate(hits))
    client = anthropic.Anthropic(api_key=get_api_key())
    msg = client.messages.create(
        model=CLAUDE_MODEL,
        max_tokens=900,
        system=(
            "You are a z/OS systems programmer assistant helping with a 3.1 to 3.2 migration. "
            "Answer ONLY from the provided context. If the context lacks the answer, say so and "
            "tell the user which IBM source (PSP bucket, HOLDDATA, FIXCAT, IBM Documentation) to check. "
            "Never invent PTF, APAR or USERMOD numbers. Cite context items like [1]."
        ),
        messages=[{"role": "user", "content": f"Context:\n{ctx}\n\nQuestion: {question}"}],
    )
    return "".join(b.text for b in msg.content if getattr(b, "type", "") == "text")


# --------------------------------------------------------------------------------------
# ML: cluster fixes
# --------------------------------------------------------------------------------------
def cluster_inventory(df: pd.DataFrame, n_clusters: int) -> pd.DataFrame:
    if len(df) < max(n_clusters, 3):
        out = df.copy()
        out["Cluster"] = 0
        return out
    texts = (df["Type"] + " " + df["Component/FMID"] + " " + df["Description"]).tolist()
    X = TfidfVectorizer(stop_words="english", sublinear_tf=True).fit_transform(texts)
    comps = max(2, min(20, X.shape[1] - 1, len(df) - 1))
    Z = TruncatedSVD(n_components=comps, random_state=42).fit_transform(X)
    labels = KMeans(n_clusters=n_clusters, n_init=10, random_state=42).fit_predict(Z)
    out = df.copy()
    out["Cluster"] = labels
    return out


# --------------------------------------------------------------------------------------
# Excel export
# --------------------------------------------------------------------------------------
def build_excel(inv: pd.DataFrame) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        summary = (
            inv.groupby(["Type", "Phase"]).size().reset_index(name="Count")
            if not inv.empty
            else pd.DataFrame(columns=["Type", "Phase", "Count"])
        )
        summary.to_excel(xw, "Summary", index=False)
        inv.to_excel(xw, "All Fixes", index=False)
        for typ, sheet in [("PTF", "PTF"), ("USERMOD", "USERMOD"), ("APAR", "APAR"), ("FIXCAT", "FIXCAT"), ("HOLD", "HOLDDATA")]:
            inv[inv["Type"] == typ].to_excel(xw, sheet, index=False)
        inv[~inv["Type"].isin(["PTF", "USERMOD", "APAR", "FIXCAT", "HOLD"])].to_excel(xw, "Other Fixes", index=False)
        SOURCES.to_excel(xw, "Sources", index=False)
        pd.DataFrame(
            {
                "Note": [
                    f"Generated {datetime.now():%Y-%m-%d %H:%M} by z/OS 3.2 Migration Fix Dashboard",
                    "Rows with Verified=False are checklist placeholders, not confirmed fixes.",
                    "Confirm every PTF/APAR/USERMOD against HOLDDATA, FIXCAT and the PSP bucket for your target zones.",
                ]
            }
        ).to_excel(xw, "Notes", index=False)

        head_fill = PatternFill("solid", fgColor="1F4E78")
        for ws in xw.book.worksheets:
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = head_fill
                c.alignment = Alignment(vertical="center")
            ws.freeze_panes = "A2"
            if ws.max_row > 1 and ws.max_column > 0:
                ws.auto_filter.ref = ws.dimensions
            for col in ws.columns:
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(12, width + 2), 70)
    return buf.getvalue()


# --------------------------------------------------------------------------------------
# UI
# --------------------------------------------------------------------------------------
st.title("🖥️ z/OS 3.1 → 3.2 Migration Fix Dashboard")
st.caption("PTFs · USERMODs · APARs · FIXCATs · RAG assistant · ML clustering · Excel export")

with st.sidebar:
    st.header("Settings")
    st.text_input(
        "Anthropic API key (optional)",
        type="password",
        key="api_key_input",
        help="Better: set ANTHROPIC_API_KEY in .streamlit/secrets.toml or as an environment variable.",
    )
    st.caption("Without a key the assistant still works in retrieval-only mode.")
    st.divider()
    usermod_regex = st.text_input(
        "USERMOD ID regex",
        value=r"\b(?:UM|SM|ZZ|LM)\d{5}\b",
        help="Sites name USERMODs differently. Adjust to your naming standard.",
    )
    default_phase = st.selectbox("Default phase for imported rows", ["Pre-migration (3.1)", "Coexistence", "Post-migration (3.2)"])
    st.divider()
    inv_all = st.session_state.inventory
    st.metric("Fixes tracked", len(inv_all))
    st.metric("Verified", int(inv_all["Verified"].sum()))

tab_over, tab_inv, tab_imp, tab_rag, tab_ml, tab_src, tab_exp = st.tabs(
    ["📊 Overview", "🗂️ Fix Inventory", "📥 Import", "🤖 RAG Assistant", "🧠 ML Clusters", "🌐 Sources", "📤 Excel Export"]
)

# ---- Overview
with tab_over:
    inv = st.session_state.inventory
    c1, c2, c3, c4 = st.columns(4)
    c1.metric("PTFs", int((inv["Type"] == "PTF").sum()))
    c2.metric("APARs", int((inv["Type"] == "APAR").sum()))
    c3.metric("USERMODs", int((inv["Type"] == "USERMOD").sum()))
    c4.metric("FIXCAT / Other", int((~inv["Type"].isin(["PTF", "APAR", "USERMOD"])).sum()))
    if not inv.empty:
        l, r = st.columns(2)
        by_type = inv.groupby("Type").size().reset_index(name="Count")
        l.plotly_chart(px.bar(by_type, x="Type", y="Count", color="Type", title="Fixes by type"), use_container_width=True)
        r.plotly_chart(px.pie(inv, names="Phase", title="Fixes by migration phase", hole=0.45), use_container_width=True)
        st.plotly_chart(
            px.treemap(inv, path=["Phase", "Type", "Priority"], title="Phase → Type → Priority"),
            use_container_width=True,
        )
    st.info(
        "The starter rows are a checklist, not confirmed fix numbers. Import your REPORT MISSINGFIX output, "
        "HOLDDATA extracts or PSP text on the **Import** tab to populate real PTF/APAR/USERMOD IDs."
    )

# ---- Inventory
with tab_inv:
    inv = st.session_state.inventory
    f1, f2, f3, f4 = st.columns(4)
    types = f1.multiselect("Type", sorted(inv["Type"].unique()), default=sorted(inv["Type"].unique()))
    phases = f2.multiselect("Phase", sorted(inv["Phase"].unique()), default=sorted(inv["Phase"].unique()))
    prios = f3.multiselect("Priority", ["High", "Medium", "Low"], default=["High", "Medium", "Low"])
    q = f4.text_input("Search")
    view = inv[inv["Type"].isin(types) & inv["Phase"].isin(phases) & inv["Priority"].isin(prios)]
    if q:
        mask = view.apply(lambda row: q.lower() in " ".join(map(str, row.values)).lower(), axis=1)
        view = view[mask]
    st.caption("Edit cells directly; add rows with the + at the bottom. Click 'Save edits' to keep changes.")
    edited = st.data_editor(
        view,
        num_rows="dynamic",
        use_container_width=True,
        height=460,
        column_config={
            "Type": st.column_config.SelectboxColumn(options=["PTF", "USERMOD", "APAR", "FIXCAT", "HOLD", "OTHER"]),
            "Priority": st.column_config.SelectboxColumn(options=["High", "Medium", "Low"]),
            "Phase": st.column_config.SelectboxColumn(options=["Pre-migration (3.1)", "Coexistence", "Post-migration (3.2)"]),
            "Verified": st.column_config.CheckboxColumn(),
        },
        key="editor",
    )
    cc1, cc2 = st.columns(2)
    if cc1.button("💾 Save edits"):
        rest = inv.drop(index=view.index)
        st.session_state.inventory = pd.concat([rest, edited], ignore_index=True)
        st.success("Saved.")
        st.rerun()
    if cc2.button("🗑️ Reset to starter checklist"):
        st.session_state.inventory = SEED.copy()
        st.rerun()

# ---- Import
with tab_imp:
    st.subheader("1) Upload SMP/E output, HOLDDATA extracts, PSP text")
    files = st.file_uploader("Files (.txt, .csv, .xlsx, .log)", accept_multiple_files=True)
    if files and st.button("Extract fixes from files"):
        total = 0
        for f in files:
            text = read_upload(f)
            add_doc(f.name, text)
            total += add_rows(extract_fixes(text, f.name, default_phase, usermod_regex))
        st.success(f"Added {total} new rows and indexed {len(files)} file(s) for the RAG assistant.")

    st.subheader("2) Paste text")
    pasted = st.text_area("Paste PSP bucket, REPORT MISSINGFIX output, HOLDDATA lines…", height=180)
    if st.button("Extract from pasted text") and pasted.strip():
        add_doc("Pasted text", pasted)
        n = add_rows(extract_fixes(pasted, "Pasted text", default_phase, usermod_regex))
        st.success(f"Added {n} new rows.")

    st.subheader("3) Fetch a web page")
    url = st.text_input("URL (IBM Docs, PSP page, support technote…)")
    if st.button("Fetch & extract") and url.strip():
        try:
            text = fetch_url_text(url.strip())
            add_doc(url, text)
            n = add_rows(extract_fixes(text, url, default_phase, usermod_regex))
            st.success(f"Fetched {len(text):,} characters; added {n} new rows and indexed the page.")
        except Exception as e:
            st.error(f"Could not fetch the page: {e}. Some IBM pages need a browser session; paste the text instead.")

# ---- RAG
with tab_rag:
    st.subheader("Ask about your migration fixes")
    st.caption("Retrieval uses TF-IDF over your inventory, source list and anything you imported.")
    question = st.text_input("Question", placeholder="Which fixes are needed on 3.1 for coexistence with 3.2?")
    k = st.slider("Passages to retrieve", 3, 12, 6)
    if question:
        hits = retrieve(question, k)
        if not hits:
            st.warning("Nothing relevant found. Import more material on the Import tab.")
        else:
            if get_api_key():
                with st.spinner("Asking Claude…"):
                    try:
                        st.markdown(ask_claude(question, hits))
                    except Exception as e:
                        st.error(f"Claude call failed: {e}")
            else:
                st.info("No API key set: showing retrieved passages only.")
            with st.expander("Retrieved context", expanded=not get_api_key()):
                for i, (h, s) in enumerate(hits, 1):
                    st.markdown(f"**[{i}] {h['title']}** · similarity {s:.2f}")
                    st.write(h["text"][:600])

# ---- ML
with tab_ml:
    st.subheader("Group related fixes (TF-IDF → SVD → KMeans)")
    inv = st.session_state.inventory
    if len(inv) < 4:
        st.info("Import more fixes to enable clustering.")
    else:
        n_clusters = st.slider("Number of clusters", 2, min(10, len(inv) - 1), min(4, len(inv) - 1))
        clustered = cluster_inventory(inv, n_clusters)
        st.dataframe(clustered.sort_values("Cluster"), use_container_width=True, height=380)
        st.plotly_chart(
            px.histogram(clustered, x="Cluster", color="Type", title="Cluster composition"),
            use_container_width=True,
        )
        if st.button("Use cluster labels in export"):
            st.session_state.inventory = clustered.copy()
            st.success("Cluster column added to the inventory.")

# ---- Sources
with tab_src:
    st.subheader("Where to get z/OS 3.2 product and fix details")
    st.caption("Deep links on IBM sites move often. Use the search hint if a page has changed.")
    for _, r in SOURCES.iterrows():
        a, b = st.columns([3, 1])
        a.markdown(f"**{r['Source']}**  \n{r['What you get / search hint']}  \n_Feeds: {r['Feeds']}_")
        b.link_button("Open", r["URL"])
    if st.button("Check which sites respond"):
        res = []
        for _, r in SOURCES.drop_duplicates("URL").iterrows():
            try:
                code = requests.head(r["URL"], timeout=8, allow_redirects=True).status_code
            except Exception as e:
                code = type(e).__name__
            res.append({"URL": r["URL"], "Status": code})
        st.dataframe(pd.DataFrame(res), use_container_width=True)

# ---- Export
with tab_exp:
    st.subheader("Excel workbook")
    st.write("Sheets: Summary, All Fixes, PTF, USERMOD, APAR, FIXCAT, HOLDDATA, Other Fixes, Sources, Notes.")
    data = build_excel(st.session_state.inventory)
    st.download_button(
        "⬇️ Download zOS_3.2_Migration_Fixes.xlsx",
        data=data,
        file_name="zOS_3.2_Migration_Fixes.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
